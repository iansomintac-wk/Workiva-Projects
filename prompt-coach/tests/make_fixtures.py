"""Generate real Office/PDF fixtures for tests/filetest.js (python-docx, openpyxl, python-pptx, reportlab, Pillow; LibreOffice for the legacy .doc)."""
import csv, os, subprocess
import docx, openpyxl, pptx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw
out = os.path.join(os.path.dirname(__file__), 'fixtures'); os.makedirs(out, exist_ok=True); os.chdir(out)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Summary'
ws.append(['Q3 FY26 Income Statement Flux']); ws.append([]); ws.append(['Account', 'Q2 FY26', 'Q3 FY26', 'Change', 'Change %'])
for a, x, y in [('Revenue', 412.3, 455.9), ('Cost of revenue', 180.2, 201.7), ('SG&A', 95.4, 112.8), ('R&D', 60.1, 61.0)]: ws.append([a, x, y, round(y - x, 1), round((y - x) / x * 100, 1)])
d = wb.create_sheet('Detail'); d.append(['Entity', 'Account', 'Amount', 'Period'])
for e in ['US', 'UK', 'DE']: d.append([e, 'Revenue', 100.0, 'Q3 FY26'])
wb.save('Q3 Flux.xlsx')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Controls'; ws.append(['Control ID', 'Description', 'Owner', 'Frequency', 'Key'])
ws.append(['REV-01', 'Revenue cutoff review', 'Controller', 'Monthly', 'Y']); ws.append(['REV-02', 'Contract approval', '', 'Per event', 'Y']); wb.save('RCM_FY26.xlsx')
doc = docx.Document(); doc.add_heading('Access Control Policy', 0); doc.add_heading('1. Purpose', 1); doc.add_paragraph('This policy sets requirements for user access reviews in FY2026.')
doc.add_heading('2. Quarterly access reviews', 1); doc.add_paragraph('Control owners must review privileged access every quarter (Q1 FY26 to Q4 FY26).'); doc.save('Access Policy.docx')
pr = pptx.Presentation()
for title, body in [('Q3 FY26 Results', 'Revenue up 10.6% vs Q2 FY26'), ('Scope 3 Emissions', 'FY2025 vs FY2024 by category'), ('Outlook', 'FY2027 guidance')]:
    s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = title; s.placeholders[1].text = body
pr.save('Board Deck Q3.pptx')
c = canvas.Canvas('GHG Inventory 2025.pdf', pagesize=letter); c.drawString(72, 720, 'GHG Inventory Report FY2025'); c.drawString(72, 700, 'Scope 1: 12,400 tCO2e (FY2024: 13,100)'); c.showPage(); c.drawString(72, 720, 'Methodology: GHG Protocol.'); c.save()
img = Image.new('RGB', (800, 400), 'white'); ImageDraw.Draw(img).text((40, 40), 'scanned page', fill='black'); img.save('scan.png')
c = canvas.Canvas('Scanned Invoice.pdf', pagesize=letter); c.drawImage('scan.png', 72, 400, width=400, height=200); c.save(); os.remove('scan.png')
with open('Supplier Data.csv', 'w', newline='') as f: w = csv.writer(f); w.writerow(['Supplier', 'FY2025 tCO2e', 'Method']); w.writerow(['Acme', '1200', 'spend-based'])
subprocess.run(['soffice', '--headless', '--convert-to', 'doc', 'Access Policy.docx'], capture_output=True)
print('fixtures written to', out)
